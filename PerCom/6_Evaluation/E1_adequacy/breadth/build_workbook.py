from pathlib import Path
import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.utils import get_column_letter

ROOT=Path(__file__).resolve().parent
corpus=pd.read_csv(ROOT/'corpus_100.csv').fillna('')
screen=pd.read_csv(ROOT/'screening_log_150.csv').fillna('')
depth=corpus[(corpus.depth_status=='completed_seed') | (corpus.depth_status=='proposed_depth_addition')].copy()
depth['_order']=depth.apply(lambda r: int(r.prior_case_id[1:].split('-')[0]) if r.depth_status=='completed_seed' else 100+int(float(r.depth_rank)),axis=1)
depth=depth.sort_values('_order').drop(columns=['_order'])
amb=corpus[corpus.screen_status!='IN_SCOPE'].copy()

wb=Workbook(); ws=wb.active; ws.title='Summary'
navy='17365D'; blue='D9EAF7'; pale='EAF2F8'; white='FFFFFF'; red='FCE4D6'; green='E2F0D9'; gray='E7E6E6'
ws['A1']='E1 Timeline IR adequacy — 100-item breadth corpus'; ws['A1'].font=Font(size=16,bold=True,color=white); ws['A1'].fill=PatternFill('solid',fgColor=navy)
ws.merge_cells('A1:F1'); ws['A2']='Version'; ws['B2']='e1-corpus-v0.1-preaudit'; ws['A3']='Repository baseline'; ws['B3']='paper @ 0788969e5d313415276a6cf89151aca8cce7c047'; ws.merge_cells('B3:F3')
ws['A5']='Metric'; ws['B5']='Value'; ws['C5']='Interpretation'
summary=[
 ('Retained corpus','=COUNTA(Corpus_100!A2:A101)','Fixed breadth corpus'),
 ('Existing completed seed','=COUNTIF(Corpus_100!B2:B101,"seed")','Existing 12/12, 41/41 exact histories'),
 ('New candidates','=COUNTIF(Corpus_100!B2:B101,"new")','Collected before new encoding'),
 ('In scope','=COUNTIF(Corpus_100!L2:L101,"IN_SCOPE")','Breadth denominator candidate'),
 ('Ambiguous','=COUNTIF(Corpus_100!L2:L101,"AMBIGUOUS")','Retained; resolve before depth IR'),
 ('Out of scope','=COUNTIF(Corpus_100!L2:L101,"OUT_OF_SCOPE")','Retained as transparent boundary record'),
 ('Unmatched','=COUNTIF(Corpus_100!L2:L101,"UNMATCHED")','Catalog/backend separation'),
 ('Depth total','=COUNTA(Depth_20!A2:A21)','12 completed + 8 proposed'),
 ('Screened candidates','=COUNTA(Screening_150!A2:A151)','Includes non-retained log'),
]
for i,row in enumerate(summary,6):
 for j,v in enumerate(row,1): ws.cell(i,j,v)
ws['A17']='Design rules'; ws['A17'].font=Font(bold=True,color=white); ws['A17'].fill=PatternFill('solid',fgColor=navy); ws.merge_cells('A17:F17')
rules=[
 'Source-diverse corpus: official 25, research 26, elicited 24, community 25 (C15 reclassified after audit; no rebalancing).',
 'Existing 12 remain the seed cohort; 12 was a Stage A work unit without a statistical sample-size rationale.',
 'E1 semantic adequacy uses human audit plus reference-runner traces. Explorer is auxiliary and cannot change E1 adequacy.',
 'rb_preliminary is machine-assisted triage only. The author screens and codes manually; no second coder or kappa is used. Final labels remain blank until the author enters them.',
 'No new Timeline IR is written until source, interpretation, assumptions, histories, and expected ACTION traces are frozen.',
]
for i,t in enumerate(rules,18): ws.cell(i,1,u'• '+t); ws.merge_cells(start_row=i,start_column=1,end_row=i,end_column=6); ws.cell(i,1).alignment=Alignment(wrap_text=True,vertical='top')
ws.column_dimensions['A'].width=28; ws.column_dimensions['B'].width=24; ws.column_dimensions['C'].width=62
ws.freeze_panes='A5'


def add_df(name,df,tabname):
    sh=wb.create_sheet(name)
    for c,col in enumerate(df.columns,1): sh.cell(1,c,col)
    for r_idx,row in enumerate(df.itertuples(index=False,name=None),2):
        for c_idx,val in enumerate(row,1):
            cell=sh.cell(r_idx,c_idx,val)
            if df.columns[c_idx-1]=='source_url' and val:
                cell.hyperlink=str(val); cell.style='Hyperlink'
    last_col=get_column_letter(len(df.columns)); last_row=len(df)+1
    tab=Table(displayName=tabname,ref=f'A1:{last_col}{last_row}')
    tab.tableStyleInfo=TableStyleInfo(name='TableStyleMedium2',showRowStripes=True,showFirstColumn=False,showLastColumn=False)
    sh.add_table(tab); sh.freeze_panes='A2'; sh.auto_filter.ref=f'A1:{last_col}{last_row}'
    for cell in sh[1]:
        cell.font=Font(bold=True,color=white); cell.fill=PatternFill('solid',fgColor=navy); cell.alignment=Alignment(wrap_text=True,vertical='center')
    for col_i,col in enumerate(df.columns,1):
        vals=[str(col)]+[str(v) for v in df[col].head(100)]
        maxlen=max(len(v) for v in vals)
        if col in {'original_text','source_context','screen_reason','depth_reason','ambiguity_for_user','catalog_note','notes'}: width=42
        elif col in {'source_title'}: width=34
        elif col in {'source_url'}: width=28
        else: width=min(max(maxlen+2,10),24)
        sh.column_dimensions[get_column_letter(col_i)].width=width
    for row in sh.iter_rows(min_row=2,max_row=last_row):
        for cell in row: cell.alignment=Alignment(vertical='top',wrap_text=True)
    for r in range(2,last_row+1): sh.row_dimensions[r].height=45
    return sh

add_df('Corpus_100',corpus,'Corpus100')
add_df('Depth_20',depth,'Depth20')
add_df('Ambiguity_Queue',amb,'AmbiguityQueue')
add_df('Screening_150',screen,'Screening150')

code=wb.create_sheet('Codebook')
code_rows=[
 ('Field / code','Meaning'),
 ('IN_SCOPE','Reactive-temporal smart-home behavior with fixable assumptions.'),
 ('AMBIGUOUS','Behavior-defining source ambiguity; no IR before user decision.'),
 ('OUT_OF_SCOPE','Real source item outside the E1 smart-home automation unit.'),
 ('UNMATCHED','Clear requirement blocked by current catalog/backend, not automatically an IR failure.'),
 ('R1–R10 / B1–B5','Use definitions in repository E1 README.md.'),
 ('L-ACCUM','Provisional internal-accumulation boundary hypothesis.'),
 ('rb_preliminary','Machine-assisted first pass; never report as adjudicated coding.'),
 ('text_form','Whether text is verbatim, source-extracted, or researcher-normalized.'),
 ('completed_seed','One of the existing 12 depth cases.'),
 ('proposed_depth_addition','One of eight new candidates; freeze before encoding.'),
]
for r,row in enumerate(code_rows,1):
 for c,v in enumerate(row,1): code.cell(r,c,v)
code.column_dimensions['A'].width=30; code.column_dimensions['B'].width=100
for cell in code[1]: cell.font=Font(bold=True,color=white); cell.fill=PatternFill('solid',fgColor=navy)
for row in code.iter_rows():
 for cell in row: cell.alignment=Alignment(wrap_text=True,vertical='top')
code.freeze_panes='A2'

for sh in wb.worksheets:
    sh.sheet_view.showGridLines=False
    sh.sheet_properties.pageSetUpPr.fitToPage=True
    sh.page_setup.fitToWidth=1; sh.page_setup.fitToHeight=0; sh.page_layout='landscape'

out=ROOT/'E1_CORPUS_WORKBOOK.xlsx'; wb.save(out)
# Structural validation
chk=load_workbook(out,data_only=False)
assert chk.sheetnames==['Summary','Corpus_100','Depth_20','Ambiguity_Queue','Screening_150','Codebook']
assert chk['Corpus_100'].max_row==101
assert chk['Depth_20'].max_row==21
assert chk['Screening_150'].max_row==151
print(out)
