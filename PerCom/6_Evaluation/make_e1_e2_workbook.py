"""E1·E2 결과 엑셀 (논문 Evaluation 에 쓰는 값만).

~/temp/bin/python make_e1_e2_workbook.py  ->  E1_E2_results.xlsx

기록된 결과만 읽는다(재실행 없음): E1_adequacy/E1_SUMMARY.md, E1_adequacy/breadth/corpus_100.csv,
E2_fidelity/runs/e2_run.timer-binding.jsonl(.gz), E2_fidelity/pairs/*.json, handoff_timer/e2_population.json.
E2 수치는 make_e2_results.py 와 같은 규칙으로 세고, 원고 값과 다르면 멈춘다.
"""
import collections
import csv
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE / "E2_fidelity"))
import make_e2_results as e2  # noqa: E402

CODES = {
    "R1": "즉시 반응", "R2": "지연 후 호출", "R3": "지속 조건 + 끊기면 다시 셈", "R4": "사건 한 번 / 다시 무장",
    "R5": "저장값이 나중 호출에 흐름", "R6": "순차·분기", "R7": "조건까지 반복", "R8": "고정 횟수 반복",
    "R9": "시계 앵커·시간대", "R10": "정해진 시각 vs 사건 경쟁",
    "B1": "진행 중 취소·재시작", "B2": "독립 흐름·인스턴스", "B3": "이벤트 기억(look-back)", "B4": "가변 간격 반복",
    "B5": "중첩 반복",
}
SOURCES = [("official", "공식 문서·예제"), ("research", "연구 논문"), ("elicited", "연구 참가자 작성"),
           ("community", "커뮤니티 요청")]
DEPTH_NOTE = {
    "C05": "감사에서 인코딩 수정: 반복 사이 간격 0",
    "C07": "감사에서 인코딩 수정: 반복 사이 간격 0",
    "C09": "AutoTap 안전 속성을 연구자가 자동화로 바꿈",
    "C19": "감사에서 인코딩 수정: 이미 있던 사람을 도착으로 잘못 봄",
    "E1-072": "감사에서 인코딩 수정: 고정 시각 대신 플랫폼 일광 입력",
    "E1-092": "해석 변경으로 다시 씀: 서로 독립인 두 자동화로 나눔",
    "E1-099": "이력 추가로 다시 씀: 시한과 겹친 움직임 사건",
    "E1-034": "해석 변경으로 다시 씀: 오븐 확인 반복. 'never' 문장을 자동화 정책으로 읽음",
    "E1-028": "'never' 문장을 자동화 정책으로 읽음",
}
PROBES = [
    ("P1", "B3 이벤트 기억", "Brackenbury et al. CHI 2019",
     "IF Sally enters the bedroom AND the sun sets WITHIN 2 hours THEN turn on the bedroom lights.",
     "4/4", "표현됨", "거절: 입력 범위가 없는 관측값"),
    ("P2", "B3 이벤트 기억", "HA Community 363863",
     "Notify with a camera capture on motion only if the door opened within the previous x minutes.",
     "4/4", "표현됨", "거절: 두 실행 시점 값을 비교하는 조건"),
    ("P3", "B4 가변 간격", "HA Community 541232",
     "Turn on a GPIO at a repeatable interval configurable via an input number.",
     "3/3", "표현됨 (1단위 해상도)", "거절: 두 실행 시점 값을 비교하는 조건"),
]
CAUSE = {"C07": "반복·타이머 중첩 (시간 예산 초과)", "E1-099": "사건 수에 따라 늘어나는 시한 (거절)"}
KIND = {"correct": "올바른 구현", "fault": "오류 주입", "llm": "LLM 생성"}

HEAD_FILL = PatternFill("solid", fgColor="DDE4EE")
SECTION_FONT = Font(bold=True, size=12)


def sheet(wb, title, header, rows, widths):
    ws = wb.create_sheet(title)
    ws.append(header)
    for c in ws[1]:
        c.font, c.fill = Font(bold=True), HEAD_FILL
        c.alignment = Alignment(wrap_text=True, vertical="center")
    for r in rows:
        ws.append(list(r))
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = ws.dimensions
    return ws


def e1_depth():
    rows = []
    for line in (HERE / "E1_adequacy" / "E1_SUMMARY.md").read_text().splitlines():
        cells = [c.strip() for c in line.strip("|").split("|")]
        if len(cells) == 7 and cells[0].isdigit():
            rows.append(cells)
    assert len(rows) == 20
    return rows


def main():
    corpus = list(csv.DictReader(open(HERE / "E1_adequacy/breadth/corpus_100.csv", encoding="utf-8-sig")))
    by_id = {r["corpus_id"]: r for r in corpus}
    wb = Workbook()
    wb.remove(wb.active)

    # ---- E2 rows (same counting rules as make_e2_results.py)
    fin = e2.load("e2_run.timer-binding.jsonl")
    src = {}
    for f in ("pairs/e1_pairs.json", "pairs/sample_388_pairs.json"):
        src.update({q["pair_id"]: q for q in e2.json.loads((e2.HERE / f).read_text())["pairs"]})
    rs = list(fin.values())
    dec = [r for r in rs if e2.decided(e2.verdict(r))]
    obs = {r["pair_id"] for r in rs if r["kind"] == "fault" and (r["reference"]["outcome"] == "REF-DIVERGE" or
           (r["explorer"].get("witness_on_reference") or {}).get("status") == "diverge")}
    agr = collections.Counter(r["agreement"] for r in dec)
    llm = [r for r in rs if r["kind"] == "llm"]
    hand = [r for r in rs if r["kind"] != "llm"]
    reqs = collections.defaultdict(list)
    for r in hand:
        reqs[src[r["pair_id"]]["base_case"]].append(r)
    full = [b for b, g in reqs.items() if all(e2.decided(e2.verdict(r)) for r in g)]
    und = collections.Counter(src[r["pair_id"]]["base_case"] for r in rs if not e2.decided(e2.verdict(r)))
    before = sum(e2.decided((r.get("explorer_binding_final") or {}).get("verdict")) for r in rs)
    changed = sum(e2.decided((r.get("explorer_binding_final") or {}).get("verdict")) and
                  r["explorer_binding_final"]["verdict"] != e2.verdict(r) for r in rs)
    orig, supp = e2.history_counts()
    confirmed = sum(agr[a] for a in e2.CONFIRMED)
    obs_rows = [fin[p] for p in obs]
    got = (len(rs), len(dec), confirmed, sum(a.startswith("DIVERGE-WITNESS") for a in agr.elements()),
           sum("FALSE" in r["agreement"] for r in rs), len(obs), agr["AGREE-EQUIV-ON-CHECKED"], agr["AGREE-DIVERGE"],
           agr[e2.CONFIRMED[2]], len(dec) - before, changed)
    want = (140, 130, 129, 1, 0, 75, 55, 66, 8, 26, 0)
    assert got == want, (got, want)

    # ---- 요약
    ws = wb.create_sheet("요약")
    ws.column_dimensions["A"].width, ws.column_dimensions["B"].width = 58, 16
    ws.column_dimensions["C"].width = 60

    def sec(title):
        ws.append([])
        ws.append([title])
        ws.cell(ws.max_row, 1).font = SECTION_FONT

    def kv(label, value, note=""):
        ws.append([label, value, note])
        ws.cell(ws.max_row, 2).alignment = Alignment(horizontal="right")

    ws.append(["E1·E2 결과 요약 (논문 Evaluation 에 쓰는 값)"])
    ws.cell(1, 1).font = Font(bold=True, size=14)
    ws.append(["생성: PerCom/6_Evaluation/make_e1_e2_workbook.py (기록된 결과만 읽음)"])
    sec("E1 코퍼스 (외부 요구 100개)")
    cnt = collections.Counter(r["screen_status"] for r in corpus)
    kv("범위 안 (IN_SCOPE)", cnt["IN_SCOPE"])
    kv("애매함 (AMBIGUOUS)", cnt["AMBIGUOUS"])
    kv("스마트홈 자동화 단위 밖 (OUT_OF_SCOPE)", cnt["OUT_OF_SCOPE"])
    kv("범위 안 요구에서 나온 경계 후보 (B1/B2/B3/B4/B5)", "6/5/2/0/1", "B4 가변 간격 반복은 나오지 않음")
    sec("E1 깊이 사례 20건")
    kv("표현됨", "20/20", "범위가 붙은 판정 3건: E1-092, E1-095, E1-028")
    kv("입력 이력 정확 일치", "53/53")
    kv("감사에서 고친 인코딩 / 해석·이력 변경으로 다시 쓴 인코딩", "4 / 3")
    kv("추가 8건 중 Explorer 거절", "3", "E1-095, E1-099, E1-028 (표현 판정에는 안 들어감)")
    sec("E1 경계 probe 3건 (20건과 별도)")
    kv("표현됨 / 정확 일치", "3/3 / 11/11 이력")
    kv("Explorer 거절", "3/3", "P2·P3 두 실행 시점 값 비교 조건, P1 입력 범위 없는 관측값")
    sec("E2 설정")
    kv("프로그램 쌍", len(rs), f"올바른 {sum(r['kind'] == 'correct' for r in rs)} + 오류 "
       f"{sum(r['kind'] == 'fault' for r in rs)} (12 계열) + LLM 생성 {len(llm)}")
    kv("정답기 입력 이력", f"{orig + supp:,}", "논문: about 49,000")
    kv("Explorer 예산 (쌍마다)", "120 s", "400,000 상태, 2,000,000 전이")
    kv("최적화 후 추가로 판정한 쌍", len(dec) - before, f"기존 판정 변화 {changed} (논문에는 26과 '변화 없음'만)")
    sec("E2 표 — Fidelity")
    kv("판정함 (EQUIV 또는 DIVERGE)", len(dec))
    kv("  정답기가 확인", confirmed, f"이력에서 EQUIV {agr['AGREE-EQUIV-ON-CHECKED']} + 이력에서 DIVERGE "
       f"{agr['AGREE-DIVERGE']} + 반례 재생 {agr[e2.CONFIRMED[2]]}")
    kv("  정답기가 JoI 를 실행 못함", 1, "C24_003/llm: 미초기화 변수 + 1")
    kv("  정답기와 어긋남", 0)
    kv("관찰 가능한 차이가 있는 오류 쌍", len(obs))
    kv("  DIVERGE / EQUIV / 판정 없음", f"{sum(e2.verdict(r) == 'DIVERGE' for r in obs_rows)} / "
       f"{sum(e2.verdict(r) == 'EQUIV' for r in obs_rows)} / "
       f"{sum(not e2.decided(e2.verdict(r)) for r in obs_rows)}")
    sec("E2 표 — Coverage")
    kv("LLM 생성 후보 판정", f"{sum(e2.decided(e2.verdict(r)) for r in llm)}/{len(llm)}")
    kv("직접 만든 쌍 판정", f"{sum(e2.decided(e2.verdict(r)) for r in hand)}/{len(hand)}")
    kv("직접 만든 요구 전부 판정", f"{len(full)}/{len(reqs)}")
    for b, cause in CAUSE.items():
        kv(f"판정 없음: {cause}", und[b], b)

    # ---- E1 코퍼스
    rows = []
    for key, name in SOURCES:
        g = [r for r in corpus if r["source_stratum"] == key]
        c = collections.Counter(r["screen_status"] for r in g)
        rows.append([name, c["IN_SCOPE"], c["AMBIGUOUS"], c["OUT_OF_SCOPE"], len(g)])
    rows.append(["합계", cnt["IN_SCOPE"], cnt["AMBIGUOUS"], cnt["OUT_OF_SCOPE"], len(corpus)])
    ws = sheet(wb, "E1 코퍼스", ["출처 유형", "범위 안", "애매함", "범위 밖", "합계"], rows, [22, 10, 10, 10, 10])
    ws.cell(ws.max_row, 1).font = Font(bold=True)
    inscope = [r for r in corpus if r["screen_status"] == "IN_SCOPE"]
    code_n = collections.Counter(c.strip() for r in inscope for c in r["rb_adjudicated"].split(",") if c.strip())
    ws.append([])
    ws.append(["코드", "뜻", f"범위 안 {len(inscope)}건 중"])
    for c in ws[ws.max_row]:
        c.font, c.fill = Font(bold=True), HEAD_FILL
    for code, meaning in CODES.items():
        ws.append([code, meaning, code_n[code]])
    ws.column_dimensions["B"].width = 26
    ws.auto_filter.ref = None

    # ---- E1 깊이 20건
    rows = []
    for n, case, cid, cohort, label, hist, explorer in e1_depth():
        r = by_id[cid]
        exact = hist.split(" / ")[1]
        rows.append([int(n), case, cid, dict(SOURCES)[r["source_stratum"]],
                     "처음 12건" if cohort.startswith("Stage A") else "추가 8건",
                     r["original_text"], r["rb_adjudicated"], exact, label, explorer, DEPTH_NOTE.get(case, "")])
    sheet(wb, "E1 깊이 20건", ["#", "사례", "코퍼스 ID", "출처", "묶음", "원문 요구", "행동 코드", "이력 정확 일치",
                              "최종 판정", "Explorer (참고, 판정에 안 들어감)", "비고"],
          rows, [5, 9, 10, 14, 11, 60, 16, 10, 30, 26, 38])

    # ---- E1 경계 probe
    sheet(wb, "E1 경계 probe", ["ID", "경계 후보", "출처", "요구 (요약)", "이력 정확 일치", "표현 판정", "Explorer"],
          PROBES, [6, 16, 24, 60, 10, 20, 34])

    # ---- E2 쌍 140
    def check(r):
        a, ref = r["agreement"], r["reference"]["outcome"]
        if a == "AGREE-EQUIV-ON-CHECKED":
            return "확인", "이력에서 차이 없음"
        if a == "AGREE-DIVERGE":
            return "확인", "이력에서 차이 있음"
        if a == e2.CONFIRMED[2]:
            why = "이력이 차이 나는 입력에 닿지 못함" if ref == "REF-EQUIV-CHECKED" else "JoI 가 이력에 없는 기기 값을 읽음"
            return "확인", f"반례 재생으로 확인 ({why})"
        if a.startswith("DIVERGE-WITNESS"):
            return "확인 불가", "정답기가 JoI 실행 못함 (미초기화 변수 + 1)"
        return "판정 없음", CAUSE[src[r["pair_id"]]["base_case"]]

    order = {"correct": 0, "fault": 1, "llm": 2}
    rows = []
    for r in sorted(rs, key=lambda r: (r["kind"] == "llm", src[r["pair_id"]]["base_case"], order[r["kind"]],
                                       r["pair_id"])):
        q = src[r["pair_id"]]
        res, detail = check(r)
        observable = ("예" if r["pair_id"] in obs else "아니오") if r["kind"] == "fault" else ""
        rows.append([r["pair_id"], q["base_case"], KIND[r["kind"]], r["family"] or "",
                     q.get("command_eng") if r["kind"] == "llm" else (q["description"] if r["kind"] == "fault" else ""),
                     e2.verdict(r), res, detail, observable])
    sheet(wb, "E2 쌍 140", ["쌍", "요구", "종류", "오류 계열", "오류 내용 / LLM 요청", "Explorer 판정", "정답기 대조",
                           "대조 근거", "관찰 가능한 차이 (오류 쌍)"],
          rows, [22, 10, 11, 16, 50, 12, 11, 40, 12])

    out = HERE / "E1_E2_results.xlsx"
    wb.save(out)
    print(out)


if __name__ == "__main__":
    main()
